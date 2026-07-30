from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from rich.console import Console

from .sample_qc import flags_path, output_dir, summary_path
from .sample_qc_ranger import evaluation_root


console = Console()


def report_path(cfg: dict) -> Path:
    return output_dir(cfg) / "GEDI_红树林林下地形样本筛选试验报告.docx"


def _set_font(run, name: str = "Microsoft YaHei", size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size, color in [("Heading 1", 15, "1F4E79"), ("Heading 2", 12, "1F4E79")]:
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)


def _paragraph(document: Document, text: str, *, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    _set_font(run, bold=bold)


def _table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        _shade(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            _set_font(run, size=9, bold=True)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            for run in cells[index].paragraphs[0].runs:
                _set_font(run, size=8.6)
    document.add_paragraph()


def _add_figure(document: Document, path: Path, caption: str) -> None:
    if not path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Cm(15.0))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_paragraph.add_run(caption)
    _set_font(run, size=9)


def _retention_rows(retention: pd.DataFrame, reference_column: str) -> list[list[str]]:
    values: list[list[str]] = []
    for row in retention.itertuples(index=False):
        base = int(row.base_qa_rows)
        provisional = int(row.provisional_screened_rows)
        repeat = int(getattr(row, f"{reference_column}_rows"))
        values.append(
            [
                str(row.REG_CODE),
                f"{base:,}",
                f"{int(row.range_candidate_rows):,}",
                f"{provisional:,}",
                f"{repeat:,}",
                f"{provisional / base * 100:.2f}%" if base else "-",
            ]
        )
    return values


def run(cfg: dict) -> Path:
    """创建可读的中文 Word 报告；报告只复述已生成 QC 指标，不编造外部验证结果。"""
    root = output_dir(cfg)
    flags = flags_path(cfg)
    audit_path = summary_path(cfg)
    retention_path = root / "qc_retention_by_region.csv"
    if not flags.exists() or not audit_path.exists() or not retention_path.exists():
        raise FileNotFoundError("缺少 QC 审计结果。请先运行 prepare-sample-qc。")
    audit = json.loads(audit_path.read_text(encoding="utf-8"))
    retention = pd.read_csv(retention_path)
    reference_threshold = float(cfg.get("sample_qc_experiment", {}).get("repeat_reference_iqr_threshold_m", 2.0))
    reference_column = f"repeat_iqr_le_{str(reference_threshold).replace('.', '_')}m"
    if f"{reference_column}_rows" not in retention.columns:
        raise ValueError(
            "QC 留存审计缺少配置指定的高置信代理列: "
            f"{reference_column}_rows。请检查 sample_qc_experiment 配置并重新运行 prepare-sample-qc。"
        )
    target = report_path(cfg)
    document = Document()
    _style_document(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("GEDI 红树林林下地形样本筛选试验报告")
    _set_font(run, size=19, bold=True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("EGM2008 聚合标签的内部稳定性、空间一致性与可预测性对照")
    _set_font(run, size=10.5)
    document.add_paragraph()

    document.add_heading("1. 结论边界", level=1)
    _paragraph(document, "本报告的目标是比较 GEDI 聚合高程标签的内部稳定性、局地空间一致性及其对 AlphaEarth 特征的可预测性。当前没有独立 LiDAR 或 RTK 参考数据，因此任何 RMSE、MAE、Bias 和 R2 均只能解释为对 GEDI 聚合标签的随机内部对照，不能表述为真实林下地形精度。")
    _paragraph(document, "特别地，产品图件的色标范围不是训练筛选阈值；极端高程样本占比低也不能证明任一阈值正确。EGM2008 正高用于统一垂直基准，但不等同于局地平均海平面或潮位基准。")

    document.add_heading("2. 数据与试验设计", level=1)
    _table(
        document,
        ["项目", "本次设置"],
        [
            ["输入样本", f"EGM2008 聚合 Parquet，共 {int(audit['rows_output']):,} 个 10 m 像元"],
            ["基础样本", "已在上游执行 quality_flag=1、degrade_flag=0、有效 elev_lowestmode 和 EGM2008 有限值检查"],
            ["随机划分", f"固定种子 {audit.get('split_seed', cfg.get('sample_qc_experiment', {}).get('split_seed', 42))}；在候选筛选前冻结 train/test"],
            ["空间单元", f"H3 {audit['h3_resolution']} 级；至少 {audit['h3_min_samples']} 个样本才计算局地一致性"],
            ["范围对照", f"{audit['range_candidate_m'][0]} 至 {audit['range_candidate_m'][1]} m，仅作为敏感性候选"],
        ],
    )
    _table(
        document,
        ["样本层", "规则", "用途"],
        [
            ["base_qa", "上游 GEDI QA 与有效 EGM2008 标签", "现有可用主样本"],
            ["range_candidate", "保留 [-20, 50] m", "范围敏感性对照，不是生态绝对范围"],
            ["spatial_consistency_candidate", "H3 局地 MAD 审计；仅标记绝对残差 > 5 m 且稳健 Z > 6 的候选异常", "局地一致性对照"],
            ["repeat_high_confidence", "elev_count >= 2 且 elev_iqr <= 1/2/5 m", "较稳定标签的代理测试集，不作为全球主训练集"],
            ["provisional_screened", "range_candidate 且非局地候选异常", "后续模型对照候选集"],
        ],
    )

    document.add_heading("3. 样本审计结果", level=1)
    _paragraph(document, "下表展示各区域的候选规则留存量。多重访低 IQR 样本规模通常远小于基础样本，因此它们应被视为标签稳定性代理集，而不是适合单独支撑全球训练的主样本。")
    _table(
        document,
        ["区域", "base_qa", "范围候选", "暂定筛选", f"重访高置信(IQR<={reference_threshold:g}m)", "暂定筛选留存率"],
        _retention_rows(retention, reference_column),
    )
    figures = root / "figures"
    _add_figure(document, figures / "01_qc_retention_by_region.png", "图 1. MEOW-14 各区候选规则的样本留存量。")
    _add_figure(document, figures / "02_qc_global_composition.png", "图 2. 全局候选规则的样本规模；不代表外部精度。")

    document.add_heading("4. 固定参数 ranger 内部对照", level=1)
    model_root = evaluation_root(cfg)
    overall_path = model_root / "qc_candidate_model_metrics_overall.csv"
    if overall_path.exists():
        overall = pd.read_csv(overall_path)
        _paragraph(document, "所有候选模型采用相同的 64 个 AlphaEarth 特征和固定随机森林参数。比较仅检验不同训练标签集合的可预测性，不能据此证明样本更加接近真实地形。")
        _table(
            document,
            ["训练规则", "测试集", "覆盖区域", "样本数", "宏平均 RMSE", "宏平均 MAE", "宏平均 Bias", "宏平均 R2"],
            [
                [str(row.candidate), str(row.test_scope), str(int(row.regions)), f"{int(row.n):,}", f"{row.rmse_macro:.3f}", f"{row.mae_macro:.3f}", f"{row.bias_macro:.3f}", f"{row.r2_macro:.3f}"]
                for row in overall.itertuples(index=False)
            ],
        )
        _add_figure(document, model_root / "figures" / "01_qc_candidate_internal_rmse.png", "图 3. 固定参数 ranger 的内部 RMSE 对照。")
    else:
        _paragraph(document, "固定参数 ranger 候选对照尚未运行。本报告保留该章节结构；运行 evaluate-sample-qc-candidates 后重新生成报告即可补入模型指标和图件。")

    document.add_heading("5. GEDI 原始质量字段 pilot", level=1)
    inventory = root / "gedi_qc_band_inventory.json"
    if inventory.exists():
        content = json.loads(inventory.read_text(encoding="utf-8"))
        _paragraph(document, "运行时确认可用的 GEDI 质量字段为：" + "、".join(content.get("selected_fields", [])) + "。这些字段先在 14 个固定 H3 小块中进行诊断，不直接应用于全量筛选。")
    else:
        _paragraph(document, "尚未进行 GEDI quality band 只读检查。后续应先运行 inspect-gedi-qc-bands，确认产品字段和数据字典含义，再提交 14 个小块 pilot。")
    _paragraph(document, "pilot 将比较 sensitivity >= 0.90、sensitivity >= 0.95、可解释的 surface/mode 信息及其组合。若某字段缺失、语义不适合地面质量，或造成明显 MEOW 区域覆盖偏差，则不将其升级为全球硬筛选规则。")

    document.add_heading("6. 文献依据与解释规则", level=1)
    _table(
        document,
        ["文献", "可支持的结论", "不能外推的结论"],
        [
            ["Wang et al. (2026), RFDTM, doi:10.5194/essd-2026-356", "红树林林下地面点应通过信号、几何、云/噪声和空间一致性等多约束处理。", "其新西兰 ICESat-2/SAR 流程和图件范围不能直接变成全球 GEDI 绝对高程阈值。"],
            ["Huang et al. (2024), Forests 15, 2064, doi:10.3390/f15122064", "林下高程控制点提取需要多准则过滤与独立验证。", "ICESat-2 ATL08 的局地阈值不能直接替代 GEDI L2A 的全球规则。"],
            ["Pronk et al. (2024), Remote Sensing 16, 2259, doi:10.3390/rs16132259", "星载激光地形质量应以独立高精度参考数据评价。", "内部一致性或模型 holdout 不能替代外部地形验证。"],
            ["Zhu et al. (2023), Remote Sensing 15, 4969, doi:10.3390/rs15204969", "地形、坡度、植被结构和观测条件会改变星载激光地形/冠层反演误差。", "非红树林短植被结论不能直接量化全球红树林阈值。"],
        ],
    )
    _paragraph(document, "因此，本试验优先输出可审计的多证据 QC 标签，而不是用单一全局绝对高程范围删样。任何候选规则若使某一 MEOW 区留存率低于基础样本的 50%，或未在稳定性代理集上表现出一致收益，均不推荐升级为全球硬筛选。")

    document.add_heading("7. 可复现流程与下一步", level=1)
    _paragraph(document, "建议顺序：1) prepare-sample-qc；2) inspect-gedi-qc-bands；3) export-gedi-qc-pilot --submit；4) 在确认 pilot 后决定是否补提取全量质量字段；5) prepare-sample-qc-model-inputs；6) evaluate-sample-qc-candidates；7) create-sample-qc-report。")
    _paragraph(document, "当获得 LiDAR/RTK 后，应统一到 EGM2008 或可追溯的共同垂直基准，并以独立站点或空间块比较候选规则的 RMSE、MAE、Bias 和区域覆盖。届时才可以选择生产筛选规则并报告真实地形精度。")

    document.add_heading("参考文献", level=1)
    for reference in [
        "Huang, J., Wang, Y. & Yu, Y. Multi-Criteria Filtration and Extraction Strategy for Understory Elevation Control Points Using ICESat-2 ATL08 Product. Forests 15, 2064 (2024). doi:10.3390/f15122064.",
        "Pronk, M., Eleveld, M. A. & Ledoux, H. Assessing Vertical Accuracy and Spatial Coverage of ICESat-2 and GEDI Spaceborne Lidar for Creating Global Terrain Models. Remote Sensing 16, 2259 (2024). doi:10.3390/rs16132259.",
        "Wang, Y. et al. RFDTM: A national-scale and wall-to-wall 30 m resolution mangrove sub-canopy topography dataset for New Zealand derived from ICESat-2 ATLAS and multi-band SAR. Earth System Science Data Discussions preprint (2026). doi:10.5194/essd-2026-356.",
        "Zhu, X. et al. Evaluation and Comparison of ICESat-2 and GEDI Data for Terrain and Canopy Height Retrievals in Short-Stature Vegetation. Remote Sensing 15, 4969 (2023). doi:10.3390/rs15204969.",
    ]:
        _paragraph(document, reference)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    console.print(f"[green]中文 Word 报告已生成: {target}[/green]")
    return target
